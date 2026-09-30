import io
import struct
import zipfile

import docx

from app.docs.extract import decode_para_text, extract_text, iter_para_text


def record(tag: int, payload: bytes, level: int = 0) -> bytes:
    size = len(payload)
    if size >= 0xFFF:
        return struct.pack("<I", tag | (level << 10) | (0xFFF << 20)) + struct.pack("<I", size) + payload
    return struct.pack("<I", tag | (level << 10) | (size << 20)) + payload


def para(text: str) -> bytes:
    return text.encode("utf-16-le")


def test_decode_para_text_skips_controls():
    ext_ctrl = struct.pack("<H", 11) + b"\x00" * 14  # 확장 제어문자(표 등) 16바이트
    payload = para("사업") + ext_ctrl + para("계획서") + struct.pack("<H", 13)
    assert decode_para_text(payload) == "사업계획서"


def test_iter_para_text_records_and_long_size():
    long_text = "가" * 3000  # 6000바이트 → 확장 크기 필드
    raw = record(66, b"\x00" * 10) + record(67, para("1. 기술개발 개요")) + record(67, para(long_text))
    out = list(iter_para_text(raw))
    assert out == ["1. 기술개발 개요", long_text]


def test_hwpx_extraction():
    xml = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<hs:sec xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section" xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph">'
        "<hp:p><hp:run><hp:t>1. 기술개발 목표</hp:t></hp:run></hp:p>"
        "<hp:p><hp:run><hp:t>※ 정량 목표를 </hp:t></hp:run><hp:run><hp:t>작성</hp:t></hp:run></hp:p>"
        "</hs:sec>"
    )
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("Contents/section0.xml", xml)
    text = extract_text("a.hwpx", buf.getvalue())
    assert text.splitlines() == ["1. 기술개발 목표", "※ 정량 목표를 작성"]


def test_docx_extraction_with_table():
    d = docx.Document()
    d.add_paragraph("평가표")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "평가항목", "배점"
    t.cell(1, 0).text, t.cell(1, 1).text = "기술성", "30"
    buf = io.BytesIO()
    d.save(buf)
    text = extract_text("a.docx", buf.getvalue())
    assert "평가표" in text and "기술성 | 30" in text


def test_txt_cp949():
    assert extract_text("a.txt", "공고문".encode("cp949")) == "공고문"
