"""샘플 기업으로 전체 흐름(LLM 없이) 실행: 매칭 → 양식 파싱 → 인터뷰 → 초안 → DOCX → 점검."""
import io

import docx

from app.db.models import Application, PlanTemplate
from app.db.seed import seed
from app.manage.alerts import build_checklist, document_alerts
from app.matching.report import DISCLAIMER, run_matching, to_markdown
from app.review.checklist import ReviewInput, rule_checks, virtual_scoring
from app.writer import planner
from app.writer.export import export_docx, plain_text_for_hwp
from app.writer.template_parser import parse_template

TEMPLATE = """1. 기술개발 목표
※ 최종 목표를 정량적 성능지표로 제시
2. 기술개발 방법
※ 단계별 추진절차와 업무분담을 작성
3. 사업화 계획
※ 목표 시장규모와 5개년 매출 계획을 작성"""
EVAL = "기술개발 목표 | 30\n기술개발 방법 | 30\n사업화 계획 | 40"


def test_full_pipeline(session):
    company = seed(session)
    item = company.items[0]

    rows = run_matching(session, company.id, item.id, use_llm=False)
    eligible = [r for r in rows if r.eligible]
    assert len(eligible) == 2  # 부처 R&D + 충남 시제품, 서울 여성기업은 불충족
    md = to_markdown(rows, company.name, item.name)
    assert "① 부처(중앙정부)" in md and "② 지자체·지역기관" in md and "요약표" in md and DISCLAIMER in md
    assert "[R&D]" in md and "[비R&D]" in md
    rnd = next(r for r in eligible if r.announcement.is_rnd)
    assert rnd.score_total > next(r for r in eligible if not r.announcement.is_rnd).score_total - 30

    sections, eval_items, how = parse_template(TEMPLATE, EVAL)
    assert how == "rule" and len(sections) == 3 and sections[2]["max_score"] == 40
    tpl = PlanTemplate(name="테스트양식", year=2026, sections=sections, eval_items=eval_items)
    session.add(tpl)
    session.flush()
    app = Application(company_id=company.id, item_id=item.id, announcement_id=rnd.announcement.id, template_id=tpl.id, title="테스트 계획서")
    session.add(app)
    session.flush()

    questions = planner.generate_questions(session, app, use_llm=False)
    assert questions and all(q["section_key"] in {"s1", "s2", "s3"} for q in questions)
    assert not any("선행연구" in q["question"] for q in questions)  # 이미 선행연구가 있으므로 묻지 않음
    app.interview = [{**q, "answer": "답변 예시" if i == 0 else ""} for i, q in enumerate(app.interview)]
    app.extras = {
        "market": {"tam": 1e12, "sam_ratio": 10, "som_ratio": 5, "tam_source": "한국에너지공단 2025", "sam_basis": "제조업 비중", "som_basis": "점유 목표"},
        "forecast": {"start_year": 2027, "rows": [{"B": 1e8 * (i + 1), "D": 4e9, "export": 0} for i in range(5)]},
    }
    session.flush()

    for s in sections:
        planner.draft_section(session, app, s["key"], use_llm=False)
    session.refresh(app)
    assert len(app.drafts) == 3
    goal = next(d for d in app.drafts if d.section_key == "s1")
    assert "피크 전력 예측 정확도" in goal.content  # 성능지표 표 반영
    biz = next(d for d in app.drafts if d.section_key == "s3")
    assert "TAM(전체시장)" in biz.content and "총매출(A=B+D)" in biz.content
    assert planner.missing_items(goal.content)  # 오프라인 초안은 [확인 필요] 자리표시자 포함

    data = export_docx(app)
    doc = docx.Document(io.BytesIO(data))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert headings == [s["title"] for s in sections]
    assert len(doc.tables) >= 2
    assert "|" not in plain_text_for_hwp(biz.content).splitlines()[-1]

    inp = ReviewInput(
        sections=[{"key": d.section_key, "title": d.section_title, "content": d.content, "eval_item": next(s["eval_item"] for s in sections if s["key"] == d.section_key), "max_score": 30} for d in app.drafts],
        eval_items=eval_items,
        indicator_names=[i.name for i in item.indicators],
    )
    findings = rule_checks(inp)
    checks = {f.check for f in findings}
    assert "확인 필요 잔존" in checks and "성능지표–절차 미연계" in checks
    scoring = virtual_scoring(inp, findings, use_llm=False)
    assert scoring["method"] == "rule" and len(scoring["items"]) == 3
    assert 0 < scoring["total"] <= 100

    app.announcement = rnd.announcement
    checklist = build_checklist(session, app)
    by_name = {c["doc_type"]: c for c in checklist}
    assert by_name["기업부설연구소 인정서"]["vault_valid"] and not by_name["국세납세증명서"]["in_vault"]
    assert any(a.title == "벤처기업확인서" for a in document_alerts(session))
