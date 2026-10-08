"""워드 → 한글 변환 테스트: 샘플 DOCX를 만들어 HWPX/HWP로 바꾸고 다시 열어 내용을 확인한다."""
import io
import struct
import zlib

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt, RGBColor
from hwpx import HwpxDocument

from app.docs.docx_to_hwp import convert_docx


def _png(w=4, h=3) -> bytes:
    raw = b"".join(b"\x00" + b"\xff\x00\x00" * w for _ in range(h))

    def chunk(tag, data):
        return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", zlib.crc32(tag + data))

    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def make_sample_docx() -> bytes:
    d = Document()
    d.add_heading("사업계획서", level=1)
    p = d.add_paragraph("기업명: ")
    r = p.add_run("에너지솔루션(주)")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(0x1F, 0x5F, 0xBF)
    d.add_paragraph("가운데 정렬 문단").alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.add_paragraph("첫째 항목", style="List Number")
    d.add_paragraph("둘째 항목", style="List Number")
    d.add_paragraph("글머리 항목", style="List Bullet")
    t = d.add_table(rows=3, cols=3)
    t.style = "Table Grid"
    t.cell(0, 0).text = "구분"
    t.cell(0, 1).merge(t.cell(0, 2)).text = "연구개발비"
    t.cell(1, 0).merge(t.cell(2, 0)).text = "1차년도"
    t.cell(1, 1).text = "정부지원금"
    t.cell(1, 2).text = "100"
    t.cell(2, 1).text = "기관부담금"
    t.cell(2, 2).text = "20"
    d.add_picture(io.BytesIO(_png()), width=Mm(40))
    d.add_page_break()
    d.add_paragraph("두 번째 쪽")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


@pytest.fixture
def sample_docx() -> bytes:
    return make_sample_docx()


def test_docx_to_hwpx_keeps_text_table_image(sample_docx):
    res = convert_docx(sample_docx, "hwpx")
    assert res.data[:2] == b"PK"
    assert res.tables == 1 and res.images == 1
    doc = HwpxDocument.open(io.BytesIO(res.data))
    text = doc.text.plain()
    for s in ["사업계획서", "에너지솔루션(주)", "1. 첫째 항목", "2. 둘째 항목", "• 글머리 항목", "연구개발비", "기관부담금", "두 번째 쪽"]:
        assert s in text, s
    table = [t for p in doc.paragraphs for t in p.tables][0]
    assert table.cell(0, 1).span == (1, 2)  # 가로 병합
    assert table.cell(1, 0).span == (2, 1)  # 세로 병합
    assert table.cell(1, 0).text == "1차년도"


def test_docx_to_hwp_binary(sample_docx):
    res = convert_docx(sample_docx, "hwp")
    assert res.data[:8] == bytes.fromhex("D0CF11E0A1B11AE1")  # HWP 5.0 = OLE 복합 파일
    doc = HwpxDocument.open(io.BytesIO(res.data))
    assert "연구개발비" in doc.text.plain()


def test_bad_format():
    with pytest.raises(ValueError):
        convert_docx(b"", "pdf")
