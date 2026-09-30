"""DOCX 내보내기(양식 섹션 순서) + 섹션별 텍스트(HWP 양식 붙여넣기용).

- Markdown 제목/목록/굵게/표를 DOCX 서식으로 변환한다.
- 출처 없는 수치 문장과 [확인 필요] 자리표시자는 노란색 형광 표시한다.
"""
from __future__ import annotations

import io
import re
from datetime import datetime

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.shared import Pt

from app.db.models import Application, PlanDraft
from app.writer.planner import MISSING, unsourced_sentences

BOLD = re.compile(r"\*\*(.+?)\*\*")


def _set_font(document: Document, name: str = "맑은 고딕", size: int = 10) -> None:
    style = document.styles["Normal"]
    style.font.name = name
    style.font.size = Pt(size)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), name)


def _add_runs(paragraph, text: str, flagged: set[str]) -> None:
    """굵게(**) 처리 + 형광 표시."""
    highlight_whole = any(f and f in text for f in flagged)
    pos = 0
    for m in BOLD.finditer(text):
        _add_plain(paragraph, text[pos : m.start()], highlight_whole)
        run = paragraph.add_run(m.group(1))
        run.bold = True
        if highlight_whole:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        pos = m.end()
    _add_plain(paragraph, text[pos:], highlight_whole)


def _add_plain(paragraph, text: str, highlight_whole: bool) -> None:
    if not text:
        return
    pos = 0
    for m in MISSING.finditer(text):
        run = paragraph.add_run(text[pos : m.start()])
        if highlight_whole:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        miss = paragraph.add_run(m.group(0))
        miss.font.highlight_color = WD_COLOR_INDEX.YELLOW
        miss.bold = True
        pos = m.end()
    run = paragraph.add_run(text[pos:])
    if highlight_whole:
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _table(document: Document, rows: list[str], flagged: set[str]) -> None:
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r":?-{2,}:?", c or "--") for c in r)]
    if not cells:
        return
    ncols = max(len(r) for r in cells)
    table = document.add_table(rows=len(cells), cols=ncols)
    table.style = "Table Grid"
    for i, row in enumerate(cells):
        for j in range(ncols):
            cell = table.cell(i, j)
            cell.text = ""
            _add_runs(cell.paragraphs[0], row[j] if j < len(row) else "", flagged)
            if i == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True


def markdown_to_docx(document: Document, text: str) -> None:
    flagged = set(unsourced_sentences(text))
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if line.strip().startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            _table(document, block, flagged)
            continue
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        m = re.match(r"^(#{1,6})\s+(.*)", stripped)
        if m:
            document.add_heading(m.group(2), level=min(len(m.group(1)) + 1, 6))
        elif re.match(r"^[-*•]\s+", stripped):
            _add_runs(document.add_paragraph(style="List Bullet"), re.sub(r"^[-*•]\s+", "", stripped), flagged)
        elif re.match(r"^\d+[.)]\s+", stripped):
            _add_runs(document.add_paragraph(style="List Number"), re.sub(r"^\d+[.)]\s+", "", stripped), flagged)
        elif stripped.startswith(">"):
            p = document.add_paragraph()
            run = p.add_run(stripped.lstrip("> "))
            run.italic = True
        else:
            _add_runs(document.add_paragraph(), stripped, flagged)
        i += 1


def ordered_drafts(app: Application) -> list[PlanDraft]:
    order = {s["key"]: n for n, s in enumerate(app.template.sections)} if app.template else {}
    return sorted(app.drafts, key=lambda d: order.get(d.section_key, 999))


def export_docx(app: Application) -> bytes:
    document = Document()
    _set_font(document)
    title = app.title or (app.announcement.title if app.announcement else "사업계획서")
    document.add_heading(title, level=0)
    meta = document.add_paragraph()
    meta.add_run(
        f"양식: {app.template.name if app.template else '-'} ({app.template.version if app.template else ''})  ·  "
        f"작성일: {datetime.now():%Y-%m-%d}"
    ).italic = True
    note = document.add_paragraph()
    run = note.add_run("노란색 표시: 출처 미표기 수치 문장 또는 [확인 필요] 항목 — 제출 전 보완하십시오.")
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    for draft in ordered_drafts(app):
        document.add_heading(draft.section_title or draft.section_key, level=1)
        markdown_to_docx(document, draft.content or "")
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def plain_text_for_hwp(content: str) -> str:
    """HWP 양식 붙여넣기용: Markdown 기호 제거, 표는 탭 구분."""
    out = []
    for line in content.splitlines():
        s = line.rstrip()
        if s.strip().startswith("|"):
            if re.fullmatch(r"\s*\|?(\s*:?-{2,}:?\s*\|)+\s*", s):
                continue
            out.append("\t".join(c.strip() for c in s.strip().strip("|").split("|")))
            continue
        s = re.sub(r"^#{1,6}\s+", "", s)
        s = BOLD.sub(r"\1", s)
        s = re.sub(r"^[-*]\s+", "○ ", s.strip()) if re.match(r"^\s*[-*]\s+", s) else s
        s = re.sub(r"^>\s*", "", s)
        out.append(s)
    return "\n".join(out).strip()
